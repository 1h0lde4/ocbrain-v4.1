"""Shared fixtures for the capability-foundation tests (no test collection here)."""
from __future__ import annotations

import io
import json
import zipfile
from typing import Any, Callable, Dict, List, Optional, Sequence

from core.capabilities.adapter_runtime import AdapterRuntime
from core.capabilities.capability import BaseAdapter, CapabilityResult, CapabilityType
from core.capabilities.foundation.models import ModelCompletion
from core.capabilities.foundation.wiring import register_foundation_capabilities
from core.capabilities.registry import CapabilityRegistry
from core.capabilities.resource import ResourceManager


# ── documents ───────────────────────────────────────────────────────────────

def make_docx(paragraphs: Sequence[str] = ("Revenue grew 12% in Q3.",), *,
              title: str = "Quarterly Report", table: Optional[List[List[str]]] = None,
              heading: str = "Findings") -> bytes:
    import docx
    d = docx.Document()
    d.add_heading(title, 0)
    d.add_heading(heading, 1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows: Sequence[Sequence[Any]] = (("Region", "Q1"), ("EU", 5)), *,
              sheet: str = "Sales", extra_sheets: Sequence[str] = ()) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    for name in extra_sheets:
        wb.create_sheet(name).append(["x", 1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pdf(pages_text: Sequence[str]) -> bytes:
    """Minimal single-font PDF; an empty string yields an image-only-style
    page (no text layer)."""
    n = len(pages_text)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{kids}] /Count {n} >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, t in enumerate(pages_text):
        cid = 5 + 2 * i
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Contents {cid} 0 R /Resources << /Font << /F1 3 0 R >> >> >>")
        stream = f"BT /F1 12 Tf 72 720 Td ({t}) Tj ET" if t else ""
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    return out


def rewrite_zip(data: bytes, replace: Dict[str, Callable[[bytes], bytes]]) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename in replace:
                payload = replace[item.filename](payload)
            zout.writestr(item.filename, payload)
    return out.getvalue()


def zip_bomb() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * 30_000_000)
    return out.getvalue()


def traversal_zip() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("../../evil.txt", b"x")
    return out.getvalue()


# ── model doubles ───────────────────────────────────────────────────────────

class ScriptedModel:
    """TextModel double: returns scripted replies, records every prompt."""
    model_id = "scripted"

    def __init__(self, replies: Any = "ok", *, provider: str = "fake",
                 model: str = "fake-1") -> None:
        self.replies = replies
        self.prompts: List[str] = []
        self.purposes: List[str] = []
        self._provider, self._model = provider, model
        self.calls = 0

    async def complete(self, prompt: str, *, purpose: str, trace_id: str):
        self.calls += 1
        self.prompts.append(prompt)
        self.purposes.append(purpose)
        reply = self.replies(prompt) if callable(self.replies) else self.replies
        if isinstance(reply, BaseException):
            raise reply
        if isinstance(reply, ModelCompletion):
            return reply
        return ModelCompletion(text=reply, provider=self._provider, model=self._model)


def analysis_reply(findings: Sequence[Dict[str, Any]], *, conclusions=("Done.",),
                   uncertainties=(), contradictions=()) -> str:
    return json.dumps({"findings": list(findings), "contradictions": list(contradictions),
                       "uncertainties": list(uncertainties),
                       "conclusions": list(conclusions)})


class FakeLlmAdapter(BaseAdapter):
    """LLM_COMPLETION double for the planner-fallback path."""

    def __init__(self) -> None:
        super().__init__()
        self.adapter_name = "fake-llm_completion"
        self.capability_type = CapabilityType.LLM_COMPLETION
        self.calls: List[str] = []

    async def execute(self, request, resources):
        self.calls.append(request.payload.get("subtask", ""))
        return CapabilityResult(success=True, output="llm fallback answer",
                                adapter_used=self.adapter_name)


def build_registry(model=None, **kwargs):
    """Registry with the three foundation capabilities and a real AdapterRuntime."""
    registry = CapabilityRegistry()
    register_foundation_capabilities(registry, text_model=model or ScriptedModel(), **kwargs)
    runtime = AdapterRuntime(registry, ResourceManager())
    return registry, runtime


FIXED_NONCE = lambda: "0123456789abcdef"  # noqa: E731
