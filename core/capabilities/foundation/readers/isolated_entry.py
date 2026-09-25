"""
core/capabilities/foundation/readers/isolated_entry.py — parser subprocess.

Run as ``python -I <this file> <reader_id> <opts_json>``; artifact bytes on
stdin; exactly one JSON document on stdout. This is the ONLY place the
third-party document libraries (pypdf, python-docx, openpyxl) are imported: the
host process never does. A parser bug, an infinite loop, a decompression bomb
or a memory blow-up therefore costs one throw-away process, not the runtime.

Self-contained on purpose: it imports nothing from the ``core`` package (no
config load, no event stream, no import chain an attacker could steer) and it
is started with ``-I`` (no PYTHONPATH, no user site, no cwd on sys.path). The
status strings below duplicate core.capabilities.descriptors.CapabilityStatus;
tests assert they cannot drift.

Baseline isolation applied here, before any library is imported (best effort,
defense in depth -- the boundary is the process + kernel limits, not Python):
  * RLIMIT_AS      address-space cap            RLIMIT_CPU   CPU seconds
  * RLIMIT_FSIZE=0 no file writes               RLIMIT_CORE=0 no core dumps
  * RLIMIT_NOFILE  few descriptors              RLIMIT_NPROC=0 no new processes
  * socket.socket  replaced with a raiser (Python-level egress tripwire)
The host adds: wall-clock kill, output cap, scrubbed environment, empty cwd.
NOT provided by this baseline (requires the Sandbox Fabric namespace backend,
see ADR-CAP-01 "Deferred"): network namespace, filesystem read jail, seccomp.

Nothing is executed from the document: no macros, no formulas (cached values
only), no external relationships or links are followed, no JavaScript.
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import zipfile

MALFORMED = "malformed"
UNSUPPORTED = "unsupported"
LIMIT_EXCEEDED = "limit_exceeded"
DEPENDENCY_UNAVAILABLE = "dependency_unavailable"
INVALID_REQUEST = "invalid_request"
READER_VERSION = "1.0.0"
_MAX_META = 200


class _Err(Exception):
    def __init__(self, status, reason, message=""):
        super().__init__(message or reason)
        self.status, self.reason, self.message = status, reason, message or reason


def _limits_setup(limits):
    try:
        import resource
    except ImportError:  # non-POSIX: no rlimits; the host wall-clock still applies
        return
    mem = int(limits.get("memory_mb", 1024)) * 1024 * 1024
    cpu = int(float(limits.get("timeout_sec", 20))) + 2
    for name, value in (("RLIMIT_AS", mem), ("RLIMIT_CPU", cpu),
                        ("RLIMIT_FSIZE", 0), ("RLIMIT_CORE", 0),
                        ("RLIMIT_NOFILE", 64), ("RLIMIT_NPROC", 0)):
        try:
            res = getattr(resource, name)
            resource.setrlimit(res, (value, value))
        except (ValueError, OSError, AttributeError):
            pass


def _block_network():
    import socket

    def _blocked(*a, **k):
        raise OSError("network access is disabled in the reader subprocess")
    socket.socket = _blocked  # type: ignore[assignment]
    socket.create_connection = _blocked  # type: ignore[assignment]


def _dep(name, dist, floor=""):
    import importlib.metadata
    import importlib.util
    if importlib.util.find_spec(name) is None:
        raise _Err(DEPENDENCY_UNAVAILABLE, f"{dist}_not_installed")
    if floor:
        try:
            have = tuple(int(n) for n in re.findall(r"\d+", importlib.metadata.version(dist))[:3])
        except Exception:
            raise _Err(DEPENDENCY_UNAVAILABLE, f"{dist}_version_unknown")
        need = tuple(int(n) for n in re.findall(r"\d+", floor)[:3])
        if have < need:
            raise _Err(DEPENDENCY_UNAVAILABLE,
                       f"{dist}_{'.'.join(map(str, have))}_below_security_floor_{floor}")


def _meta(d):
    out = {}
    for k, v in d.items():
        if isinstance(v, str) and v.strip():
            out[str(k)[:64]] = v.strip()[:_MAX_META]
    return out


def _paragraphs(text):
    parts = [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]
    return parts


# ── PDF ─────────────────────────────────────────────────────────────────────

def read_pdf(data, opts):
    # Security floor, kept in sync with readers/defaults.py's
    # PYPDF_SECURITY_FLOOR (this module is deliberately standalone -- see the
    # module docstring -- so it cannot import that constant). Verified via
    # direct CVE/changelog lookup 2026-09-22: covers CVE-2026-54530/-54531/
    # -54651/-59935/-59936/-84309 (crafted-PDF infinite-loop DoS, fixed
    # 6.13.0 through 6.16.0) plus 6.16.1's further, not-yet-CVE-numbered
    # iteration-limit hardening. Older versions are refused rather than
    # trusted; the wall-clock kill in isolation.py is the backstop either way.
    _dep("pypdf", "pypdf", "6.16.1")
    from pypdf import PdfReader
    limits = opts.get("limits") or {}
    max_pages = int(limits.get("max_pages", 200))
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise _Err(UNSUPPORTED, "encrypted_pdf", "encrypted PDFs are not opened")
        total = len(reader.pages)
    except _Err:
        raise
    except Exception as e:
        raise _Err(MALFORMED, "pdf_parse_error", f"{type(e).__name__}")
    warnings = []
    try:
        root = reader.trailer["/Root"]
        names = root.get("/Names") or {}
        if "/OpenAction" in root or "/AA" in root or "/JavaScript" in names:
            warnings.append({"code": "active_content_present",
                             "message": "PDF declares actions/JavaScript (not executed)",
                             "locator": ""})
        if "/EmbeddedFiles" in names:
            warnings.append({"code": "embedded_files_present",
                             "message": "PDF has embedded files (not extracted)",
                             "locator": ""})
    except Exception:
        pass
    wanted = (opts.get("selection") or {}).get("pages")
    truncated = False
    if wanted is not None:
        if (not isinstance(wanted, list) or not wanted
                or not all(isinstance(n, int) and not isinstance(n, bool) for n in wanted)):
            raise _Err(INVALID_REQUEST, "bad_pages_selector", "pages must be a list of ints")
        pages = sorted({n for n in wanted if 1 <= n <= total})[:max_pages]
        for n in sorted(set(wanted) - set(pages)):
            warnings.append({"code": "selection_not_found",
                             "message": f"page {n} does not exist", "locator": f"page={n}"})
    else:
        pages = list(range(1, min(total, max_pages) + 1))
        truncated = total > max_pages
    blocks, empty_pages, failed = [], 0, 0
    for n in pages:
        try:
            text = reader.pages[n - 1].extract_text() or ""
        except Exception:
            failed += 1
            warnings.append({"code": "page_extract_failed", "message": "",
                             "locator": f"page={n}"})
            continue
        paras = _paragraphs(text)
        if not paras:
            empty_pages += 1
            continue
        for k, para in enumerate(paras, 1):
            blocks.append({"kind": "paragraph", "text": para, "page": n,
                           "locator": f"page={n};para={k}"})
    if pages and not blocks and not failed and wanted is None:
        raise _Err(UNSUPPORTED, "no_text_layer",
                   "PDF has pages but no extractable text (image-only; OCR is not available)")
    if empty_pages:
        warnings.append({"code": "pages_without_text",
                         "message": f"{empty_pages} page(s) have no text layer",
                         "locator": ""})
    md = {}
    try:
        info = reader.metadata
        if info:
            md = _meta({"title": info.title, "author": info.author,
                        "subject": info.subject, "creator": info.creator})
    except Exception:
        pass
    return {"title": md.get("title", ""), "metadata": md, "page_count": total,
            "blocks": blocks, "tables": [], "warnings": warnings,
            "truncated": truncated, "truncation_reason": "max_pages" if truncated else "",
            "failed_pages": failed}


# ── OOXML (docx / xlsx) ─────────────────────────────────────────────────────

def _zip_preflight(data, limits):
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except zipfile.BadZipFile:
        raise _Err(MALFORMED, "bad_zip")
    if len(infos) > int(limits.get("max_zip_entries", 2000)):
        raise _Err(LIMIT_EXCEEDED, "zip_too_many_entries")
    total, ratio_cap = 0, int(limits.get("max_zip_ratio", 100))
    for i in infos:
        parts = i.filename.replace("\\", "/").split("/")
        if ("\x00" in i.filename or i.filename.startswith(("/", "\\"))
                or ".." in parts or re.match(r"^[A-Za-z]:", i.filename)):
            raise _Err(MALFORMED, "zip_path_traversal")
        if i.flag_bits & 0x1:
            raise _Err(UNSUPPORTED, "encrypted_zip_entry")
        total += i.file_size
        if i.file_size > 1_000_000 and i.compress_size > 0 \
                and i.file_size / i.compress_size > ratio_cap:
            raise _Err(LIMIT_EXCEEDED, "zip_compression_ratio")
    if total > int(limits.get("max_zip_uncompressed_bytes", 200_000_000)):
        raise _Err(LIMIT_EXCEEDED, "zip_uncompressed_size")
    return {i.filename for i in infos}


def read_ooxml(data, opts):
    limits = opts.get("limits") or {}
    names = _zip_preflight(data, limits)
    warnings = []
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        warnings.append({"code": "macros_present",
                         "message": "document contains VBA macros (never executed)",
                         "locator": ""})
    if any("/embeddings/" in n or n.lower().startswith("embeddings/") for n in names):
        warnings.append({"code": "embedded_objects_present",
                         "message": "document embeds objects (not extracted)",
                         "locator": ""})
    if "word/document.xml" in names:
        return _read_docx(data, limits, warnings)
    if "xl/workbook.xml" in names:
        return _read_xlsx(data, opts, limits, warnings)
    raise _Err(UNSUPPORTED, "unsupported_zip_container",
               "not a .docx or .xlsx package")


def _read_docx(data, limits, warnings):
    _dep("docx", "python-docx")
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    max_rows, max_cols = int(limits.get("max_rows", 2000)), int(limits.get("max_cols", 100))
    max_tables = int(limits.get("max_tables", 200))
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise _Err(MALFORMED, "docx_parse_error", type(e).__name__)
    blocks, tables, title, truncated = [], [], "", False
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = (p.text or "").strip()
            if not text:
                continue
            style = ""
            try:
                style = (p.style.name if p.style is not None else "") or ""
            except Exception:
                pass
            m = re.match(r"^Heading (\d)$", style)
            if m or style == "Title":
                level = int(m.group(1)) if m else 1
                if style == "Title" and not title:
                    title = text
                blocks.append({"kind": "heading", "level": level, "text": text,
                               "locator": f"block={len(blocks) + 1}"})
            elif style.lower().startswith("list"):
                blocks.append({"kind": "list_item", "text": text,
                               "locator": f"block={len(blocks) + 1}"})
            else:
                blocks.append({"kind": "paragraph", "text": text,
                               "locator": f"block={len(blocks) + 1}"})
        elif tag == "tbl":
            if len(tables) >= max_tables:
                truncated = True
                continue
            t, rows, cut = Table(child, d), [], False
            for ri, row in enumerate(t.rows):
                if ri >= max_rows:
                    cut = True
                    break
                rows.append([(c.text or "").strip() for c in row.cells[:max_cols]])
            tables.append({"rows": rows, "header_rows": 1 if len(rows) > 1 else 0,
                           "truncated": cut, "row_count": len(rows),
                           "locator": f"table={len(tables) + 1}"})
            blocks.append({"kind": "table", "table_ref": len(tables), "text": "",
                           "locator": f"table={len(tables)}"})
            if len(rows) > 1:
                warnings.append({"code": "header_row_assumed",
                                 "message": "first table row treated as header",
                                 "locator": f"table={len(tables)}"})
    md = {}
    try:
        cp = d.core_properties
        md = _meta({"title": cp.title, "author": cp.author, "subject": cp.subject})
    except Exception:
        pass
    return {"title": title or md.get("title", ""), "metadata": md, "page_count": None,
            "blocks": blocks, "tables": tables, "warnings": warnings,
            "truncated": truncated, "truncation_reason": "max_tables" if truncated else ""}


def _col_letter(n):
    out = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out or "A"


def _read_xlsx(data, opts, limits, warnings):
    _dep("openpyxl", "openpyxl")
    import openpyxl
    max_rows, max_cols = int(limits.get("max_rows", 2000)), int(limits.get("max_cols", 100))
    max_sheets = int(limits.get("max_sheets", 20))
    wanted = (opts.get("selection") or {}).get("sheets")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise _Err(MALFORMED, "xlsx_parse_error", type(e).__name__)
    blocks, tables, truncated = [], [], False
    try:
        names = list(wb.sheetnames)
        if wanted is not None:
            if not isinstance(wanted, list) or not all(isinstance(s, str) for s in wanted):
                raise _Err(INVALID_REQUEST, "bad_sheets_selector", "sheets must be a list of names")
            for missing in [s for s in wanted if s not in names]:
                warnings.append({"code": "selection_not_found",
                                 "message": f"sheet {missing!r} does not exist",
                                 "locator": f"sheet={missing}"})
            names = [n for n in names if n in wanted]
        elif len(names) > max_sheets:
            truncated = True
            names = names[:max_sheets]
        for name in names:
            ws = wb[name]
            if getattr(ws, "sheet_state", "visible") != "visible":
                warnings.append({"code": "hidden_sheet_skipped", "message": "",
                                 "locator": f"sheet={name}"})
                continue
            rows, cut = [], False
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                if ri >= max_rows:
                    cut = True
                    break
                rows.append(["" if v is None else str(v) for v in list(row)[:max_cols]])
            while rows and not any(c.strip() for c in rows[-1]):
                rows.pop()
            if not rows:
                continue
            width = max(len(r) for r in rows)
            tables.append({"rows": rows, "header_rows": 1 if len(rows) > 1 else 0,
                           "truncated": cut, "sheet": name[:128],
                           "range": f"A1:{_col_letter(width)}{len(rows)}",
                           "row_count": len(rows), "col_count": width,
                           "locator": f"sheet={name};range=A1:{_col_letter(width)}{len(rows)}"})
            blocks.append({"kind": "heading", "level": 1, "text": name[:128],
                           "locator": f"sheet={name}"})
            blocks.append({"kind": "table", "table_ref": len(tables), "text": "",
                           "locator": f"sheet={name}"})
            if cut:
                truncated = True
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return {"title": "", "metadata": {}, "page_count": None, "blocks": blocks,
            "tables": tables, "warnings": warnings, "truncated": truncated,
            "truncation_reason": "max_rows_or_sheets" if truncated else ""}


_READERS = {"pdf": read_pdf, "ooxml": read_ooxml}


def main(argv):
    try:
        reader_id, opts = argv[1], json.loads(argv[2])
        fn = _READERS[reader_id]
    except Exception:
        _emit({"error": {"status": INVALID_REQUEST, "reason": "bad_invocation"}})
        return
    limits = opts.get("limits") or {}
    _block_network()
    data = sys.stdin.buffer.read(int(limits.get("max_bytes", 10_000_000)) + 1)
    _limits_setup(limits)
    try:
        if len(data) > int(limits.get("max_bytes", 10_000_000)):
            raise _Err(LIMIT_EXCEEDED, "artifact_too_large")
        body = fn(data, opts)
        body.setdefault("reader", {"id": reader_id, "version": READER_VERSION})
        _emit({"ok": body})
    except _Err as e:
        _emit({"error": {"status": e.status, "reason": e.reason, "message": e.message[:240]}})
    except MemoryError:
        _emit({"error": {"status": LIMIT_EXCEEDED, "reason": "memory_limit"}})
    except Exception as e:  # a parser blowing up on hostile input == malformed input
        _emit({"error": {"status": MALFORMED, "reason": f"parser_exception:{type(e).__name__}"}})


def _emit(obj):
    sys.stdout.buffer.write(json.dumps(obj, ensure_ascii=True).encode("ascii"))
    sys.stdout.buffer.flush()


if __name__ == "__main__":
    main(sys.argv)
